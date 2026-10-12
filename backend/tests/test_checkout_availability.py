"""Checkout configuration gates billing without restricting safe conversion."""

import base64
import csv
import io
import json
import time
from contextlib import closing
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from reportlab.pdfgen import canvas

import statement_to_excel.api as api
from statement_to_excel.billing import (
    CONVERSION_IDEMPOTENCY_TTL_SECONDS,
    MockBillingProvider,
    PersistentLedger,
    StripeBillingProvider,
    apply_webhook,
)
from tests.test_billing_stripe import PaymentDouble, buy, paid_event, sign
from tests.test_pdf_edge_api import _failure_pdf


TEST_KEY = "sk_test_synthetic_not_a_credential"
WEBHOOK_SECRET = "whsec_synthetic_not_a_credential"
TABLES = (
    "accounts", "conversion_charges", "checkout_requests", "checkout_aliases",
    "entitlements", "webhook_events", "subscription_state",
)
UNAVAILABLE_CONFIGURATION = [
    pytest.param(("", ""), id="empty-keys"),
    pytest.param((TEST_KEY, ""), id="test-key-only"),
    pytest.param(("", WEBHOOK_SECRET), id="webhook-secret-only"),
]


@pytest.fixture
def ten_page_pdf():
    buf = io.BytesIO()
    pdf = canvas.Canvas(buf, pagesize=(612, 792))
    for day in range(1, 11):
        pdf.drawString(50, 750, "Synthetic Checking Statement")
        pdf.drawString(50, 730, "Statement period 01/01/2026 - 01/31/2026")
        pdf.drawString(50, 710, "Date Description Amount Balance")
        pdf.drawString(50, 680, f"01/{day:02d} Synthetic deposit +1.00 {1000 + day:.2f}")
        pdf.showPage()
    pdf.save()
    return buf.getvalue()


@pytest.fixture(params=UNAVAILABLE_CONFIGURATION)
def unavailable_client(request, monkeypatch, tmp_path):
    ledger = PersistentLedger(str(tmp_path / "billing.sqlite3"))
    monkeypatch.setattr(api, "LEDGER", ledger)
    monkeypatch.setattr(api, "PROVIDER", StripeBillingProvider(*request.param))
    _forbid_stripe_network(monkeypatch)
    with TestClient(api.app) as client:
        yield client, ledger


def _forbid_stripe_network(monkeypatch):
    import stripe

    def forbidden(*args, **kwargs):
        pytest.fail("No Checkout request should reach Stripe in this test")

    monkeypatch.setattr(stripe.checkout.Session, "create", forbidden)


def _forbid_conversion_billing(monkeypatch, ledger):
    def forbidden(*args, **kwargs):
        pytest.fail("Unmetered conversion must not create accounts, hash uploads or debit pages")

    for method in ("get_or_create", "charge_pages", "charge_pages_once", "conversion_charge"):
        monkeypatch.setattr(ledger, method, forbidden)
    # Replace only the API's module reference, preserving hashing used by PDF
    # libraries and the real Stripe signature verifier.
    monkeypatch.setattr(api, "hashlib", SimpleNamespace(sha256=forbidden))


def billing_counts(ledger):
    with closing(ledger.connect()) as db:
        return {name: db.execute(f"SELECT count(*) FROM {name}").fetchone()[0]
                for name in ("accounts", "conversion_charges", "checkout_requests")}


def _snapshot(ledger):
    with closing(ledger.connect()) as db:
        return {table: [dict(row) for row in db.execute(f"SELECT * FROM {table} ORDER BY 1")]
                for table in TABLES}


def _assert_rows(rows):
    assert rows == [{
        "date": f"2026-01-{day:02d}", "description": "Synthetic deposit",
        "amount": "1.00", "balance": f"{1000 + day:.2f}",
        "currency": "USD", "type": "credit",
    } for day in range(1, 11)]


def _assert_download(data, fmt):
    if fmt == "csv":
        _assert_rows(list(csv.DictReader(io.StringIO(data.decode("utf-8")))))
        return
    workbook = load_workbook(io.BytesIO(data), data_only=True)
    try:
        assert workbook.sheetnames == ["Transactions", "Summary"]
        rows = list(workbook["Transactions"].values)
        assert rows[0] == ("Date", "Description", "Amount", "Balance", "Currency", "Type")
        assert len(rows) == 11
        for day, row in enumerate(rows[1:], 1):
            assert row[0].isoformat().startswith(f"2026-01-{day:02d}")
            assert row[1:] == ("Synthetic deposit", 1, 1000 + day, "USD", "credit")
        summary = dict(workbook["Summary"].values)
        assert summary["Transactions"] == 10
        assert summary["Total credit"] == summary["Net"] == 10
    finally:
        workbook.close()


@pytest.mark.parametrize(("configuration", "expected"), [
    pytest.param(("", ""), False, id="empty-keys"),
    pytest.param((TEST_KEY, ""), False, id="test-key-only"),
    pytest.param(("", WEBHOOK_SECRET), False, id="webhook-secret-only"),
    pytest.param((TEST_KEY, "wrong-webhook-prefix"), False, id="invalid-webhook-prefix"),
    pytest.param((TEST_KEY, WEBHOOK_SECRET), True, id="both-stripe-test-prefixes"),
    pytest.param("mock", False, id="mock-is-not-configured-checkout"),
    pytest.param("named-double", False, id="stripe-name-alone-is-insufficient"),
])
def test_availability_is_consistent_on_provider_health_and_pricing(monkeypatch, tmp_path, configuration, expected):
    if configuration == "mock":
        provider = MockBillingProvider()
    elif configuration == "named-double":
        provider = SimpleNamespace(name="stripe", checkout_available=True)
    else:
        provider = StripeBillingProvider(*configuration)
        assert provider.checkout_available is expected
    ledger = PersistentLedger(str(tmp_path / "billing.sqlite3"))
    monkeypatch.setattr(api, "LEDGER", ledger)
    monkeypatch.setattr(api, "PROVIDER", provider)
    _forbid_stripe_network(monkeypatch)
    with TestClient(api.app) as client:
        assert api.checkout_available() is expected
        health = client.get("/api/health")
        pricing = client.get("/api/pricing")
        assert health.status_code == pricing.status_code == 200
        assert health.json()["checkout_available"] is expected
        assert pricing.json()["checkout_available"] is expected
        assert health.json()["billing"] == provider.name
        assert pricing.json()["plans"] and pricing.json()["credit_packs"]
        assert TEST_KEY not in health.text + pricing.text
        assert WEBHOOK_SECRET not in health.text + pricing.text
    assert _snapshot(ledger) == {table: [] for table in TABLES}


@pytest.mark.parametrize(("fmt", "preview"), [
    ("json", False), ("csv", False), ("xlsx", False), ("csv", True), ("xlsx", True),
])
def test_unavailable_ten_page_pdf_converts_twice_without_quota_or_metadata(
    unavailable_client, ten_page_pdf, monkeypatch, fmt, preview,
):
    client, ledger = unavailable_client
    _forbid_conversion_billing(monkeypatch, ledger)
    for _ in range(2):
        result = client.post("/api/convert", files={"file": ("synthetic.pdf", ten_page_pdf, "application/pdf")},
                             data={"fmt": fmt, "preview": str(preview).lower()},
                             headers={"Idempotency-Key": "reusable-unmetered-request-key"})
        assert result.status_code == 200, result.text
        assert "X-Credits-Remaining" not in result.headers
        if fmt == "json" or preview:
            body = result.json()
            _assert_rows(body["transactions"])
            assert body["summary"]["count"] == 10
            assert body["charge"] == {"metered": False, "pages": 10}
            if preview:
                assert body["filename"] == f"statement.{fmt}"
                _assert_download(base64.b64decode(body["file_b64"], validate=True), fmt)
        else:
            assert result.headers["X-Transaction-Count"] == "10"
            assert f'filename="statement.{fmt}"' in result.headers["Content-Disposition"]
            _assert_download(result.content, fmt)
        assert billing_counts(ledger) == {'accounts': 0, 'conversion_charges': 0, 'checkout_requests': 0}
        assert _snapshot(ledger) == {table: [] for table in TABLES}


def test_unavailable_account_quote_checkout_and_webhook_write_nothing(unavailable_client, monkeypatch):
    client, ledger = unavailable_client
    _forbid_conversion_billing(monkeypatch, ledger)
    account = client.get("/api/account/unknown-synthetic-account")
    assert account.status_code == 200
    assert account.json() == {
        "id": "unknown-synthetic-account", "plan": "free", "credits": 0,
        "included_pages": 3, "pages_used_this_cycle": 0,
        "included_remaining": 3, "pro": False,
    }
    quote = client.post("/api/account/unknown-synthetic-account/quote", json={"pages": 10})
    assert quote.status_code == 200
    assert quote.json() == {"metered": False, "pages": 10, "sufficient": True}
    assert client.post("/api/account/unknown-synthetic-account/quote", json={"pages": -1}).status_code == 400
    checkout = client.post("/api/checkout", json={"account_id": "unknown-synthetic-account", "sku": "business"},
                           headers={"Idempotency-Key": "unavailable-valid-checkout-key"})
    assert checkout.status_code == 503
    assert client.post("/api/stripe/webhook", content=b"{}").status_code == 503
    assert _snapshot(ledger) == {table: [] for table in TABLES}


def test_unavailable_preserves_paid_state_and_blocks_valid_signed_payment(
    unavailable_client, ten_page_pdf, monkeypatch,
):
    client, ledger = unavailable_client
    provider = PaymentDouble()
    pack = buy(ledger, provider)
    assert apply_webhook(ledger, paid_event(pack["session_id"]))["granted"]
    plan = buy(ledger, provider, "business", key="synthetic-paid-plan-request")
    assert apply_webhook(ledger, paid_event(plan["session_id"], "business", event_id="evt_paid_plan"))["granted"]
    ledger.charge_pages_once("buyer", 5, "fresh-conversion-metadata", "synthetic-fresh-fingerprint")
    ledger.charge_pages_once("buyer", 1, "expired-conversion-metadata", "synthetic-expired-fingerprint")
    with ledger.transaction() as db:
        db.execute("UPDATE conversion_charges SET created_at=? WHERE request_key=?",
                   (time.time() - CONVERSION_IDEMPOTENCY_TTL_SECONDS - 60, "expired-conversion-metadata"))
    pending = buy(ledger, provider, "pack_200", key="synthetic-pending-pack-request")
    payload = json.dumps(paid_event(pending["session_id"], "pack_200", event_id="evt_blocked_payment")).encode()
    signature = sign(payload)
    # Exercise the actual Stripe SDK verifier: this payload is valid, including
    # for the dangerous webhook-secret-only configuration.
    verified = StripeBillingProvider(TEST_KEY, WEBHOOK_SECRET).verify_event(payload, signature)
    assert verified["data"]["object"]["payment_status"] == "paid"
    if api.PROVIDER.webhook_secret:
        assert api.PROVIDER.verify_event(payload, signature)["id"] == "evt_blocked_payment"
    before = _snapshot(ledger)
    expected = {**before, "conversion_charges": [row for row in before["conversion_charges"]
                                                 if row["request_key"] != "expired-conversion-metadata"]}
    _forbid_conversion_billing(monkeypatch, ledger)
    existing_account = client.get("/api/account/buyer").json()
    assert existing_account["plan"] == "business"
    assert existing_account["pro"] and existing_account["credits"] == 50
    assert existing_account["pages_used_this_cycle"] == 6
    assert client.post("/api/account/buyer/quote", json={"pages": 10}).json() == {
        "metered": False, "pages": 10, "sufficient": True,
    }
    for _ in range(2):
        assert client.post("/api/checkout", json={"account_id": "buyer", "sku": "pack_1000"},
                           headers={"Idempotency-Key": "unavailable-checkout-stable-key"}).status_code == 503
        blocked = client.post("/api/stripe/webhook", content=payload,
                              headers={"Stripe-Signature": signature})
        assert blocked.status_code == 503
        assert _snapshot(ledger) == before
    for _ in range(2):
        response = client.post("/api/convert", files={"file": ("synthetic.pdf", ten_page_pdf, "application/pdf")},
                               data={"account_id": "buyer", "fmt": "json"},
                               headers={"Idempotency-Key": "fresh-conversion-metadata"})
        assert response.status_code == 200, response.text
        assert response.json()["charge"] == {"metered": False, "pages": 10}
        _assert_rows(response.json()["transactions"])
        assert client.get("/api/account/buyer").json() == existing_account
        assert _snapshot(ledger) == expected


@pytest.mark.parametrize(("kind", "status", "message"), [
    ("bad format", 400, "fmt must be xlsx|csv|json"),
    ("empty", 400, "empty upload"),
    ("oversized", 413, "limited to 20 MiB"),
    ("invalid", 422, "Invalid or damaged PDF"),
    ("encrypted", 422, "Password-protected PDF"),
    ("scanned", 422, "scanned/image-only"),
    ("mixed scan", 422, "mixed scanned/text PDFs"),
    ("header scan", 422, "scanned body beneath selectable headers"),
    ("multiple accounts", 422, "Multiple accounts detected"),
])
def test_unavailable_keeps_upload_validation_without_billing_writes(
    unavailable_client, ten_page_pdf, monkeypatch, kind, status, message,
):
    client, ledger = unavailable_client
    _forbid_conversion_billing(monkeypatch, ledger)
    if kind == "bad format":
        data = ten_page_pdf
    elif kind == "empty":
        data = b""
    elif kind == "oversized":
        data = b"x" * (20 * 1024 * 1024 + 1)
    else:
        data = _failure_pdf(kind).getvalue()
    response = client.post("/api/convert", files={"file": ("synthetic-invalid.pdf", data, "application/pdf")},
                           data={"fmt": "pdf" if kind == "bad format" else "json"},
                           headers={"Idempotency-Key": "synthetic-invalid-request-key"})
    assert response.status_code == status, response.text
    assert message in response.json()["detail"]
    assert _snapshot(ledger) == {table: [] for table in TABLES}


def test_configured_stripe_retains_free_allowance_and_ten_page_quota(monkeypatch, tmp_path, fixtures_dir, ten_page_pdf):
    ledger = PersistentLedger(str(tmp_path / "billing.sqlite3"))
    monkeypatch.setattr(api, "LEDGER", ledger)
    monkeypatch.setattr(api, "PROVIDER", StripeBillingProvider(TEST_KEY, WEBHOOK_SECRET))
    _forbid_stripe_network(monkeypatch)
    one_page = (fixtures_dir / "acmebank.pdf").read_bytes()
    with TestClient(api.app) as client:
        assert client.get("/api/health").json()["checkout_available"] is True
        assert client.get("/api/pricing").json()["checkout_available"] is True
        assert client.get("/api/account/configured-synthetic").json()["included_remaining"] == 3
        for used in range(1, 4):
            response = client.post("/api/convert", files={"file": ("synthetic.pdf", one_page, "application/pdf")},
                                   data={"account_id": "configured-synthetic", "fmt": "json"},
                                   headers={"Idempotency-Key": f"configured-free-request-{used}"})
            assert response.status_code == 200, response.text
            assert response.json()["charge"]["from_included"] == 1
            assert response.json()["charge"]["from_credits"] == 0
            assert response.json()["charge"].get("metered") is not False
            assert client.get("/api/account/configured-synthetic").json()["pages_used_this_cycle"] == used
        exhausted = client.post("/api/convert", files={"file": ("synthetic.pdf", one_page, "application/pdf")},
                                data={"account_id": "configured-synthetic", "fmt": "json"})
        assert exhausted.status_code == 402
        assert exhausted.json()["error"] == "insufficient_credits"
        quote = client.post("/api/account/configured-ten-pages/quote", json={"pages": 10})
        assert quote.status_code == 200
        assert quote.json()["from_included"] == 3
        assert quote.json()["from_credits"] == 7
        assert quote.json()["sufficient"] is False
        before = _snapshot(ledger)
        denied = client.post("/api/convert", files={"file": ("synthetic.pdf", ten_page_pdf, "application/pdf")},
                             data={"account_id": "configured-ten-pages", "fmt": "xlsx", "preview": "true"},
                             headers={"Idempotency-Key": "configured-ten-pages-request"})
        assert denied.status_code == 402
        assert denied.json()["quote"]["pages"] == 10
        assert _snapshot(ledger) == before
    assert billing_counts(ledger) == {"accounts": 2, "conversion_charges": 3, "checkout_requests": 0}
