import io
import json
from decimal import Decimal

from openpyxl import load_workbook

from statement_to_excel.transaction import Transaction
from statement_to_excel.reconciliation import reconcile
from statement_to_excel.writers import (
    to_csv_str,
    to_json_str,
    to_xlsx_bytes,
)


def _sample():
    return [
        Transaction(
            "2024-01-05", "PAYROLL", Decimal("3250.00"), balance=Decimal("5250.00")
        ),
        Transaction(
            "2024-01-08", "WHOLE FOODS", Decimal("-84.21"), balance=Decimal("5165.79")
        ),
    ]


def _checked(rows, totals, account_type="deposit"):
    class StatementRows(list):
        pass

    result = StatementRows(rows)
    result.reconciliation = reconcile(result, totals, account_type=account_type)
    return result


def test_csv_header_and_values():
    csv = to_csv_str(_sample())
    lines = csv.strip().splitlines()
    assert lines[0] == "date,description,amount,balance,currency,type"
    assert lines[1] == "2024-01-05,PAYROLL,3250.00,5250.00,USD,credit"
    assert lines[2].endswith("-84.21,5165.79,USD,debit")


def test_json_structure_and_exact_decimals():
    data = json.loads(to_json_str(_sample()))
    assert "transactions" in data
    assert data["transactions"][0]["amount"] == "3250.00"
    assert data["transactions"][1]["amount"] == "-84.21"
    assert data["reconciliation"]["status"] == "cannot_check"
    assert data["reconciliation"]["printed_totals"]["opening"] is None
    assert data["layout"] == "unknown"


def test_xlsx_has_three_sheets_and_numeric_cells():
    raw = to_xlsx_bytes(_sample())
    wb = load_workbook(io.BytesIO(raw))
    assert wb.sheetnames == ["Transactions", "Summary", "Checks"]
    ws = wb["Transactions"]
    # header
    assert [c.value for c in ws[1]] == [
        "Date",
        "Description",
        "Amount",
        "Balance",
        "Currency",
        "Type",
    ]
    # amount cell is a real number, not a string
    amount_cell = ws.cell(row=2, column=3).value
    assert isinstance(amount_cell, (int, float))
    assert abs(amount_cell - 3250.00) < 1e-9
    # summary nets out
    summ = wb["Summary"]
    rows = {r[0].value: r[1].value for r in summ.iter_rows(min_row=2)}
    assert rows["Transactions"] == 2
    assert abs(rows["Net"] - (3250.00 - 84.21)) < 1e-9


def test_xlsx_empty_list_still_valid():
    raw = to_xlsx_bytes([])
    wb = load_workbook(io.BytesIO(raw))
    assert "Transactions" in wb.sheetnames
    assert wb["Checks"]["B3"].value == "cannot_check"
    assert wb["Checks"]["C4"].value is None
    assert wb["Checks"]["C7"].value is None


def test_json_reports_found_totals_and_flags_without_changing_csv_columns():
    rows = _sample()
    rows[1].flags = ("multiline_description", "sign_from_balance")
    checked = _checked(rows, {"opening": Decimal("2000.00"), "closing": Decimal("5165.79")})
    checked.layout = "deposit_table"
    data = json.loads(to_json_str(checked))
    assert data["reconciliation"]["status"] == "balanced"
    assert data["layout"] == "deposit_table"
    assert data["reconciliation"]["difference"] == "0.00"
    assert data["reconciliation"]["printed_totals"] == {
        "opening": "2000.00", "money_in": None, "money_out": None, "closing": "5165.79",
    }
    assert data["reconciliation"]["row_flags"] == [{
        "row_index": 1, "flags": ["multiline_description", "sign_from_balance"],
    }]
    assert to_csv_str(checked) == to_csv_str(_sample())
    assert len(data["transactions"][1]) == 6


def test_checks_preserve_overdraft_large_amounts_and_live_transaction_references():
    rows = [
        Transaction("2024-01-01", "Synthetic overdraft", Decimal("-5.00"), Decimal("-25.00")),
        Transaction("2024-01-02", "Synthetic deposit", Decimal("1000000.01"), Decimal("999975.01")),
    ]
    checked = _checked(rows, {
        "opening": Decimal("-20.00"), "money_in": Decimal("1000000.01"),
        "money_out": Decimal("5.00"), "closing": Decimal("999975.01"),
    })
    wb = load_workbook(io.BytesIO(to_xlsx_bytes(checked)))
    sheet = wb["Checks"]
    assert sheet["B3"].value == "balanced"
    assert sheet["C4"].value == -20
    assert sheet["C5"].value == 1000000.01
    assert sheet["C7"].value == 999975.01
    assert sheet["B5"].value == '=SUMIF(Transactions!$C$2:$C$3,">0",Transactions!$C$2:$C$3)'
    assert sheet["B6"].value == '=-SUMIF(Transactions!$C$2:$C$3,"<0",Transactions!$C$2:$C$3)'
    assert sheet["B7"].data_type == "f"
    assert "B4+B5-B6" in sheet["B7"].value
    assert "ISNUMBER(C4)" in sheet["B9"].value
    assert "ISNUMBER(C7)" in sheet["B9"].value
    assert "Q2:Q3" in sheet["B9"].value
    assert wb.calculation.fullCalcOnLoad and wb.calculation.forceFullCalc


def test_card_checks_keep_categories_and_missing_printed_zeros_separate():
    rows = [
        Transaction("2024-01-01", "Synthetic payment", Decimal("125.00"), category="payment_credit"),
        Transaction("2024-01-02", "Synthetic purchase", Decimal("-40.00"), category="purchase"),
        Transaction("2024-01-03", "Synthetic fee", Decimal("-3.00"), category="fee"),
        Transaction("2024-01-04", "Synthetic interest", Decimal("-2.00"), category="interest"),
    ]
    totals = {key: Decimal(value) for key, value in {
        "previous": "100.00", "payments_credits": "125.00", "purchases": "40.00",
        "fees": "3.00", "interest": "2.00", "new": "20.00",
    }.items()}
    checked = _checked(rows, totals, "card")
    wb = load_workbook(io.BytesIO(to_xlsx_bytes(checked)))
    sheet = wb["Checks"]
    assert sheet["B3"].value == "balanced"
    assert sheet["I2"].value == "=Transactions!C2"
    assert sheet["J3"].value == "=-Transactions!C3"
    assert sheet["K4"].value == "=-Transactions!C4"
    assert sheet["L5"].value == "=-Transactions!C5"
    assert "B4-B5+B6+B7+B8" in sheet["B9"].value
    assert "ISNUMBER(C8)" in sheet["B11"].value
    assert sheet["C9"].value == 20

    del totals["interest"]
    wb = load_workbook(io.BytesIO(to_xlsx_bytes(_checked(rows, totals, "card"))))
    assert wb["Checks"]["B3"].value == "cannot_check"
    assert wb["Checks"]["C8"].value is None


def test_descending_transactions_check_running_balances_oldest_first():
    rows = [
        Transaction("2024-01-03", "Synthetic credit", Decimal("7.00"), Decimal("102.00")),
        Transaction("2024-01-02", "Synthetic debit", Decimal("-5.00"), Decimal("95.00")),
    ]
    checked = _checked(rows, {"opening": Decimal("100.00"), "closing": Decimal("102.00")})
    sheet = load_workbook(io.BytesIO(to_xlsx_bytes(checked)))["Checks"]
    assert sheet["N2"].value == 3
    assert sheet["P2"].value == "=Transactions!D3"
    assert sheet["N3"].value == 2
    assert "O2+Transactions!C2" in sheet["O3"].value


def test_unknown_card_category_cannot_silently_drop_cancelling_rows_from_checks():
    rows = [
        Transaction("2024-01-01", "Synthetic unknown payment", Decimal("50.00"), category="unknown"),
        Transaction("2024-01-02", "Synthetic unknown charge", Decimal("-50.00"), category="unknown"),
    ]
    totals = {key: Decimal("0.00") for key in ("payments_credits", "purchases", "fees", "interest")}
    totals.update(previous=Decimal("100.00"), new=Decimal("100.00"))
    checked = _checked(rows, totals, "card")
    assert checked.reconciliation.difference == 0
    assert checked.reconciliation.status == "not_balanced"
    sheet = load_workbook(io.BytesIO(to_xlsx_bytes(checked)))["Checks"]
    # Unknown metadata must retain both money movements in the live subtotal
    # checks even though their effect on the final balance cancels.
    assert sheet["I2"].value == "=Transactions!C2"
    assert sheet["J3"].value == "=-Transactions!C3"
    assert sheet["B5"].value == "=SUM(I2:I3)"
    assert sheet["B6"].value == "=SUM(J2:J3)"
    assert sheet["C5"].value == sheet["C6"].value == 0
