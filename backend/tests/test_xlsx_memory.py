"""XLSX generation must never write statement cells to temporary files."""
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from decimal import Decimal
from io import BytesIO
import tempfile

from openpyxl import load_workbook
from openpyxl.worksheet import _writer

from statement_to_excel.transaction import Transaction
from statement_to_excel.writers import to_xlsx_bytes


def _forbid_temporary_files(monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("Statement conversion must not create temporary files")

    for name in ("NamedTemporaryFile", "TemporaryFile", "mkstemp", "mkdtemp"):
        monkeypatch.setattr(tempfile, name, forbidden)
    # openpyxl imports NamedTemporaryFile directly; patch the actual call site.
    monkeypatch.setattr(_writer, "NamedTemporaryFile", forbidden)
    monkeypatch.setattr(_writer, "create_temporary_file", forbidden)


def test_xlsx_without_temporary_files_preserves_cells_and_formatting(monkeypatch):
    _forbid_temporary_files(monkeypatch)
    txns = [
        Transaction("2024-01-05", "SYNTHETIC PAYROLL", Decimal("3250.00"), Decimal("5250.00")),
        Transaction("2024-01-08", "=SYNTHETIC_LITERAL_TEXT", Decimal("-84.21"), Decimal("5165.79")),
    ]
    workbook = load_workbook(BytesIO(to_xlsx_bytes(txns)))
    assert workbook.sheetnames == ["Transactions", "Summary"]
    sheet = workbook["Transactions"]
    assert list(sheet.values) == [
        ("Date", "Description", "Amount", "Balance", "Currency", "Type"),
        (datetime(2024, 1, 5), "SYNTHETIC PAYROLL", 3250, 5250, "USD", "credit"),
        (datetime(2024, 1, 8), "=SYNTHETIC_LITERAL_TEXT", -84.21, 5165.79, "USD", "debit"),
    ]
    assert sheet["B3"].data_type == "s"
    assert sheet["C3"].data_type == "n"
    assert sheet["A3"].number_format == "yyyy-mm-dd"
    assert sheet["C3"].number_format == "#,##0.00;[Red](#,##0.00)"
    assert sheet["D3"].number_format == sheet["C3"].number_format
    assert sheet["A1"].font.bold
    assert sheet["A1"].fill.fgColor.rgb == "001F2937"
    assert sheet.freeze_panes == "A2"
    assert sheet.column_dimensions["B"].width == 48
    summary = {row[0]: row[1] for row in list(workbook["Summary"].values)[1:]}
    assert summary == {
        "Transactions": 2,
        "Total debit": 84.21,
        "Total credit": 3250,
        "Net": 3165.79,
        "Currency": "USD",
        "Start date": "2024-01-05",
        "End date": "2024-01-08",
    }


def test_parallel_xlsx_conversions_keep_rows_isolated_without_temporary_files(monkeypatch):
    _forbid_temporary_files(monkeypatch)

    def convert(index):
        transaction = Transaction("2024-01-05", f"SYNTHETIC ACCOUNT {index}", Decimal(index))
        workbook = load_workbook(BytesIO(to_xlsx_bytes([transaction])))
        return workbook["Transactions"]["B2"].value, workbook["Transactions"]["C2"].value

    with ThreadPoolExecutor(max_workers=4) as workers:
        rows = list(workers.map(convert, range(12)))
    assert rows == [(f"SYNTHETIC ACCOUNT {index}", index) for index in range(12)]
