"""A deadline must stop the work, rather than merely stop waiting for a thread."""

import asyncio
import io
import os
import tempfile
import time

import anyio
import httpx
import pytest
from fastapi import HTTPException, UploadFile

import statement_to_excel.api as api
import statement_to_excel.conversion_timeout as deadline
from statement_to_excel.billing import Ledger
from tests.deadline_support import WorkerProbe


@pytest.mark.parametrize(("phase", "fmt", "preview"), [
    ("count", "json", False), ("parse", "json", False),
    ("csv", "csv", False), ("xlsx", "xlsx", False),
    ("json", "json", False), ("base64", "xlsx", True),
    ("json", "csv", True),
    ("partial_output", "csv", False), ("after_output", "json", False),
])
def test_uncooperative_document_phase_is_killed_reaped_and_retry_succeeds(
    monkeypatch, fixtures_dir, configured_provider, phase, fmt, preview,
):
    ledger = Ledger()
    monkeypatch.setattr(api, "LEDGER", ledger)
    original_command = deadline._worker_command
    data = (fixtures_dir / "acmebank.pdf").read_bytes()
    key = "synthetic-timeout-retry-123456"

    async def exercise():
        transport = httpx.ASGITransport(app=api.app)
        async with httpx.AsyncClient(transport=transport, base_url="http://testserver") as client:
            async def request():
                return await client.post("/api/convert", files={
                    "file": ("SECRET_SYNTHETIC_FILENAME.pdf", data, "application/pdf"),
                }, data={"account_id": "deadline", "fmt": fmt, "preview": str(preview).lower()},
                   headers={"Idempotency-Key": key})

            async with WorkerProbe(monkeypatch, phase) as probe:
                started = time.monotonic()
                task = asyncio.create_task(request())
                await probe.wait_entered()
                health = await asyncio.wait_for(client.get("/api/health"), timeout=1)
                assert health.status_code == 200 and health.json()["status"] == "ok"
                assert not task.done(), "health must respond while document work is still blocked"
                response = await asyncio.wait_for(task, timeout=5)
                elapsed = time.monotonic() - started
                assert response.status_code == 504, response.text
                assert "try again" in response.json()["detail"].lower()
                assert "SECRET_SYNTHETIC" not in response.text
                assert probe.timeout * 0.75 <= elapsed <= probe.timeout + 3
                probe.assert_reaped()
                assert all(process.returncode < 0 for process in probe.processes)
                assert ledger.accounts == {} and ledger._conversion_charges == {} and ledger.log == []

            monkeypatch.setattr(deadline, "_worker_command", original_command)
            monkeypatch.setattr(deadline, "CONVERSION_TIMEOUT_SECONDS", 30)
            first = await request()
            replay = await request()
            assert first.status_code == replay.status_code == 200, first.text
            assert ledger.accounts["deadline"].pages_used_this_cycle == 1
            assert len(ledger._conversion_charges) == 1
            if fmt == "json" or preview:
                assert first.json()["summary"]["count"] == 5
                assert "reconciliation" in first.json()
                assert first.json()["charge"] == replay.json()["charge"]
            elif fmt == "xlsx":
                from openpyxl import load_workbook

                workbook = load_workbook(io.BytesIO(first.content))
                assert workbook["Transactions"].max_row == 6
                assert "Checks" in workbook.sheetnames
                workbook.close()
            else:
                assert len(first.text.splitlines()) == 6

    asyncio.run(exercise())


@pytest.mark.parametrize("phase", ["crash", "exit", "oversized_output", "malformed_output",
                                   "stdout_flood", "memory"])
def test_failed_or_unbounded_worker_protocol_is_safe_reaped_and_never_charged(
    monkeypatch, fixtures_dir, configured_provider, phase,
):
    ledger = Ledger()
    monkeypatch.setattr(api, "LEDGER", ledger)

    async def exercise():
        async with WorkerProbe(monkeypatch, phase, timeout=5) as probe:
            transport = httpx.ASGITransport(app=api.app)
            async with httpx.AsyncClient(transport=transport, base_url="http://testserver") as client:
                response = await asyncio.wait_for(client.post("/api/convert", files={
                    "file": ("synthetic.pdf", (fixtures_dir / "acmebank.pdf").read_bytes(), "application/pdf"),
                }, data={"account_id": "failed-worker", "fmt": "xlsx"},
                   headers={"Idempotency-Key": "synthetic-failed-worker-123456"}), timeout=8)
            assert response.status_code == 500, response.text
            assert "try again" in response.json()["detail"].lower()
            assert "SECRET_SYNTHETIC" not in response.text
            probe.assert_reaped()
            assert ledger.accounts == {} and ledger._conversion_charges == {} and ledger.log == []

    asyncio.run(exercise())


@pytest.mark.parametrize("cancel", ["timeout", "asyncio", "anyio"])
def test_timeout_and_cancellation_delete_forced_disk_upload_before_return(
    monkeypatch, fixtures_dir, tmp_path, configured_provider, cancel,
):
    ledger = Ledger()
    monkeypatch.setattr(api, "LEDGER", ledger)
    spool = tempfile.SpooledTemporaryFile(max_size=1, mode="w+b", dir=tmp_path)
    spool.write((fixtures_dir / "acmebank.pdf").read_bytes())
    spool.rollover()
    descriptor = spool.fileno()
    spool.seek(0)
    upload = UploadFile(file=spool, filename="synthetic-forced-spool.pdf")

    async def exercise():
        async with WorkerProbe(monkeypatch, "parse", timeout=1.5 if cancel == "timeout" else 3) as probe:
            async def request():
                return await api.convert(file=upload, account_id="forced-spool", fmt="csv",
                                         currency="USD", preview=False,
                                         idempotency_key="synthetic-forced-spool-123456")

            if cancel == "anyio":
                async with anyio.create_task_group() as requests:
                    requests.start_soon(request)
                    await probe.wait_entered()
                    requests.cancel_scope.cancel()
            else:
                task = asyncio.create_task(request())
                await probe.wait_entered()
                if cancel == "asyncio":
                    task.cancel()
                    with pytest.raises(asyncio.CancelledError):
                        await task
                else:
                    with pytest.raises(HTTPException) as failure:
                        await task
                    assert failure.value.status_code == 504
            probe.assert_reaped()
            assert spool.closed
            with pytest.raises(OSError):
                os.fstat(descriptor)
            assert list(tmp_path.iterdir()) == []
            assert ledger.accounts == {} and ledger._conversion_charges == {} and ledger.log == []

    try:
        asyncio.run(exercise())
    finally:
        spool.close()


@pytest.mark.parametrize("value", ["0", "-1", "nan", "inf", "121", "invalid"])
def test_invalid_timeout_setting_fails_closed(monkeypatch, value):
    monkeypatch.setenv("STX_CONVERSION_TIMEOUT_SECONDS", value)
    with pytest.raises(ValueError):
        deadline._deadline_setting()


def test_timeout_default_and_valid_override(monkeypatch):
    monkeypatch.delenv("STX_CONVERSION_TIMEOUT_SECONDS", raising=False)
    assert deadline._deadline_setting() == 30
    monkeypatch.setenv("STX_CONVERSION_TIMEOUT_SECONDS", "2.5")
    assert deadline._deadline_setting() == 2.5


def test_real_worker_applies_memory_cpu_and_no_core_guards_before_parsing(monkeypatch, fixtures_dir):
    async def exercise():
        async with WorkerProbe(monkeypatch, "resource_guards", timeout=5) as probe:
            artifact = await deadline.run_conversion((fixtures_dir / "acmebank.pdf").read_bytes(),
                                                     "json", "USD", False)
            assert artifact.pages == 1 and artifact.transaction_count == 5
            probe.assert_reaped()

    asyncio.run(exercise())
