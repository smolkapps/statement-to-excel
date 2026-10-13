"""Exercise printed evidence through every public preview/download path."""

import base64
import csv
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas

from statement_to_excel import api
from statement_to_excel.billing import Ledger, StripeBillingProvider


def synthetic_statement(*, closing="110.00", totals=True):
    output = io.BytesIO()
    pdf = canvas.Canvas(output, pagesize=LETTER, invariant=1)
    lines = [
        "SYNTHETIC CHECKING STATEMENT",
        "Statement period: December 1, 2023 - January 31, 2024",
        "Account ending 0000",
    ]
    if totals:
        lines += ["Opening balance 100.00", "Total deposits 20.00",
                  "Total withdrawals 10.00", f"Closing balance {closing}"]
    lines += ["Date Description Amount Balance",
              "Dec 31 OPENING BALANCE 100.00",
              "Dec 31 SYNTHETIC DEPOSIT 20.00 120.00",
              "    fictional continuation",
              "Jan 02 SYNTHETIC PURCHASE 10.00 110.00"]
    for index, line in enumerate(lines):
        pdf.drawString(40, 750 - index * 18, line)
    pdf.showPage()
    pdf.save()
    return output.getvalue()


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(api, "PROVIDER", StripeBillingProvider("", ""))
    monkeypatch.setattr(api, "LEDGER", Ledger())
    return TestClient(api.app)


@pytest.mark.parametrize("fmt", ["json", "xlsx", "csv"])
@pytest.mark.parametrize("status, closing, totals", [
    ("balanced", "110.00", True),
    ("not_balanced", "111.00", True),
    ("cannot_check", "110.00", False),
])
def test_reconciliation_in_json_and_preview(client, fmt, status, closing, totals):
    response = client.post("/api/convert", data={"fmt": fmt, "preview": "1"},
                           files={"file": ("synthetic.pdf", synthetic_statement(closing=closing, totals=totals),
                                           "application/pdf")})
    assert response.status_code == 200, response.text
    body = response.json()
    reconciliation = body["reconciliation"]
    assert reconciliation["status"] == status
    assert body["layout"] == "deposit_lines"
    assert body["transactions"][0]["date"] == "2023-12-31"
    assert body["transactions"][1]["date"] == "2024-01-02"
    assert body["transactions"][1]["amount"] == "-10.00"
    assert body["charge"] == {"metered": False, "pages": 1}
    if status == "not_balanced":
        assert reconciliation["difference"] == "-1.00"
        flags = {flag for row in reconciliation["row_flags"] for flag in row["flags"]}
        assert {"year_inferred_rollover", "sign_from_balance", "multiline_description"} <= flags
    if status == "cannot_check":
        assert reconciliation["printed_totals"]["closing"] is None
    if fmt == "csv":
        csv_rows = list(csv.reader(io.StringIO(base64.b64decode(body["file_b64"]).decode())))
        assert csv_rows[0] == ["date", "description", "amount", "balance", "currency", "type"]
        assert len(csv_rows) == 3
    elif fmt == "xlsx":
        workbook = load_workbook(io.BytesIO(base64.b64decode(body["file_b64"])))
        assert "Checks" in workbook.sheetnames
        assert any(cell.data_type == "f" for row in workbook["Checks"] for cell in row)
        if status == "not_balanced":
            assert any(cell.comment for row in workbook["Transactions"] for cell in row)
