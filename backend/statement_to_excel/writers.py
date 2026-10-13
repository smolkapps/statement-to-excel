"""Serialize a list of :class:`Transaction` to CSV / JSON / XLSX.

XLSX is the headline format (it's "statement *to excel*"): the workbook gets a
formatted ``Transactions`` sheet with real numeric/date cells plus a
``Summary`` sheet with debit/credit/net totals and a ``Checks`` sheet whose
formulas compare the editable transaction cells with printed statement figures.
"""

from __future__ import annotations

import csv
import io
import json
from datetime import date
from decimal import Decimal
from typing import Sequence

from .transaction import Transaction
from .reconciliation import reconcile

__all__ = [
    "to_csv_bytes",
    "to_json_bytes",
    "to_xlsx_bytes",
    "to_csv_str",
    "to_json_str",
]

_COLUMNS = ["date", "description", "amount", "balance", "currency", "type"]
_FLAG_NOTES = {
    "year_inferred_rollover": "The year was inferred across a December–January statement period.",
    "sign_from_balance": "The sign was derived from the change in running balance.",
    "multiline_description": "The description was joined from several printed lines.",
}


def to_csv_str(txns: Sequence[Transaction]) -> str:
    """Return a CSV string with a header row and exact decimal amounts."""
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(_COLUMNS)
    for t in txns:
        d = t.to_dict()
        writer.writerow([d[c] if d[c] is not None else "" for c in _COLUMNS])
    return buf.getvalue()


def to_csv_bytes(txns: Sequence[Transaction]) -> bytes:
    return to_csv_str(txns).encode("utf-8")


def to_json_str(txns: Sequence[Transaction]) -> str:
    """Return transactions and their statement checks with exact Decimal strings."""
    return json.dumps(
        {
            "transactions": [t.to_dict() for t in txns],
            "reconciliation": _reconciliation(txns).to_dict(),
            "layout": getattr(txns, "layout", "unknown"),
        },
        indent=2,
        ensure_ascii=False,
    )


def to_json_bytes(txns: Sequence[Transaction]) -> bytes:
    return to_json_str(txns).encode("utf-8")


def _to_float(value):
    # openpyxl wants native numbers for numeric cells; Decimal -> float is safe
    # for display (we keep exact values in CSV/JSON). Guard against None.
    if value is None:
        return None
    return float(value) if isinstance(value, Decimal) else value


def _reconciliation(txns):
    """A plain list has no printed totals and therefore cannot pass a check."""
    return getattr(txns, "reconciliation", None) or reconcile(txns, {})


def _add_checks(wb, txns, header_font, header_fill, money_fmt):
    """Write live formulas without treating absent printed figures as zero."""
    check = _reconciliation(txns)
    sheet = wb.create_sheet("Checks")
    sheet.append(["Metric", "Calculated from transactions", "Printed on statement", "Difference"])
    sheet.append(["Statement type", check.account_type])
    sheet.append(["Parser reconciliation", check.status])
    card = check.account_type == "card"
    if card:
        keys = ("previous", "payments_credits", "purchases", "fees", "interest", "new")
        labels = ("Previous balance", "Payments and credits", "Purchases", "Fees", "Interest", "New balance")
    else:
        keys = ("opening", "money_in", "money_out", "closing")
        labels = ("Opening balance", "Money in", "Money out", "Closing balance")
    for row, (key, label) in enumerate(zip(keys, labels), 4):
        sheet.cell(row, 1, label)
        sheet.cell(row, 3, _to_float(check.printed_totals[key]))
        sheet.cell(row, 4, f'=IF(AND(ISNUMBER(B{row}),ISNUMBER(C{row})),ROUND(B{row}-C{row},2),"")')

    # Start only from independently printed figures. Never infer an opening
    # balance from a transaction or substitute a missing total with zero.
    sheet["B4"] = '=IF(ISNUMBER(C4),C4,"")'
    last_transaction = max(2, len(txns) + 1)
    amount_range = f"Transactions!$C$2:$C${last_transaction}"
    if card:
        categories = {"payment_credit": 9, "purchase": 10, "fee": 11, "interest": 12}
        for column, label in enumerate(("Transaction row", "Payments and credits", "Purchases", "Fees", "Interest"), 8):
            sheet.cell(1, column, label)
        for index, transaction in enumerate(txns, 2):
            sheet.cell(index, 8, index)
            category = (transaction.category if transaction.category in categories else
                        ("payment_credit" if transaction.amount > 0 else "purchase"))
            for key, column in categories.items():
                sign = "" if key == "payment_credit" else "-"
                sheet.cell(index, column, f"={sign}Transactions!C{index}" if key == category else 0)
        for row, column in ((5, "I"), (6, "J"), (7, "K"), (8, "L")):
            sheet.cell(row, 2, f"=SUM({column}2:{column}{last_transaction})")
        sheet["B9"] = '=IF(ISNUMBER(C4),B4-B5+B6+B7+B8,"")'
        closing_row = 9
        required = [f"ISNUMBER(C{row})" for row in range(4, 10)]
    else:
        sheet["B5"] = f'=SUMIF({amount_range},">0",{amount_range})'
        sheet["B6"] = f'=-SUMIF({amount_range},"<0",{amount_range})'
        sheet["B7"] = '=IF(ISNUMBER(C4),B4+B5-B6,"")'
        closing_row = 7
        required = ["ISNUMBER(C4)", "ISNUMBER(C7)"]

    # Running balances are independently printed row figures too. Follow the
    # same date order as the pure check, including descending statements.
    ordered = sorted(enumerate(txns), key=lambda item: item[1].date)
    if len(txns) > 1 and all(a.date >= b.date for a, b in zip(txns, txns[1:])) \
            and any(a.date > b.date for a, b in zip(txns, txns[1:])):
        ordered = list(reversed(list(enumerate(txns))))
    for column, label in enumerate(("Transaction row", "Calculated running balance", "Printed running balance", "Difference"), 14):
        sheet.cell(1, column, label)
    for helper_row, (index, transaction) in enumerate(ordered, 2):
        transaction_row = index + 2
        sheet.cell(helper_row, 14, transaction_row)
        previous = "$C$4" if helper_row == 2 else f"O{helper_row - 1}"
        sign = "-" if card else "+"
        sheet.cell(helper_row, 15, f'=IF(ISNUMBER({previous}),{previous}{sign}Transactions!C{transaction_row},"")')
        if transaction.balance is not None:
            sheet.cell(helper_row, 16, f"=Transactions!D{transaction_row}")
            sheet.cell(helper_row, 17, f'=IF(AND(ISNUMBER(O{helper_row}),ISNUMBER(P{helper_row})),ROUND(O{helper_row}-P{helper_row},2),"")')

    status_row = closing_row + 2
    sheet.cell(status_row, 1, "Excel reconciliation")
    # ISNUMBER guards ensure optional/missing figures remain unknown. Numeric
    # differences include printed cashflow/card subtotals and running balances,
    # preventing opposite errors from cancelling only in the closing equation.
    differences = f"D4:D{closing_row}"
    row_differences = f"Q2:Q{last_transaction}"
    subtotal_errors = f"SUMPRODUCT(ISNUMBER({differences})*({differences}<>0))"
    row_errors = f"SUMPRODUCT(ISNUMBER({row_differences})*({row_differences}<>0))"
    sheet.cell(status_row, 2, f'=IF(AND({",".join(required)}),IF({subtotal_errors}+{row_errors}=0,"balanced","not_balanced"),"cannot_check")')
    sheet.cell(status_row + 2, 1, "Difference = calculated minus printed. Missing printed figures stay blank.")
    sheet.cell(status_row + 3, 1, "Excel differences are rounded to cents. Parser reconciliation uses exact Decimal values.")
    for row in sheet.iter_rows():
        for cell in row:
            if cell.column in (2, 3, 4, 9, 10, 11, 12, 15, 16, 17) and (cell.data_type == "f" or isinstance(cell.value, (int, float))):
                cell.number_format = money_fmt
    for cell in sheet[1]:
        if cell.value is not None:
            cell.font = header_font
            cell.fill = header_fill
    for column, width in {"A": 27, "B": 28, "C": 25, "D": 18, "H": 18, "I": 24, "J": 18, "K": 18, "L": 18, "N": 18, "O": 28, "P": 25, "Q": 18}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "B4"


def to_xlsx_bytes(txns: Sequence[Transaction]) -> bytes:
    """Build a styled .xlsx workbook in memory and return its bytes.

    Requires the ``openpyxl`` dependency (a hard dependency of this package).
    """
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet._writer import WorksheetWriter
    from openpyxl.writer.excel import ExcelWriter
    from openpyxl.workbook.properties import CalcProperties
    from zipfile import ZIP_DEFLATED, ZipFile

    wb = Workbook()
    wb.calculation = CalcProperties(fullCalcOnLoad=True, forceFullCalc=True)
    ws = wb.active
    ws.title = "Transactions"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")
    money_fmt = "#,##0.00;[Red](#,##0.00)"

    ws.append([c.capitalize() for c in _COLUMNS])
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for t in txns:
        iso = t.date
        try:
            y, m, d = (int(x) for x in iso.split("-"))
            date_cell = date(y, m, d)
        except Exception:
            date_cell = iso
        ws.append(
            [
                date_cell,
                t.description,
                _to_float(t.amount),
                _to_float(t.balance),
                t.currency,
                t.type,
            ]
        )
        # PDF text is data: openpyxl otherwise interprets a leading '=' as an
        # executable formula in the downloaded workbook.
        for column in (2, 5, 6):
            ws.cell(row=ws.max_row, column=column).data_type = "s"

    # Number formats on amount + balance columns; date format on the date column.
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd"
        row[2].number_format = money_fmt
        if row[3].value is not None:
            row[3].number_format = money_fmt

    # Column widths.
    widths = {1: 12, 2: 48, 3: 14, 4: 14, 5: 10, 6: 8}
    for idx, w in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = "A2"

    for flagged in _reconciliation(txns).row_flags:
        row = flagged["row_index"] + 2
        if 2 <= row <= ws.max_row:
            note = "Check this statement row:\n" + "\n".join(
                _FLAG_NOTES.get(flag, flag.replace("_", " ")) for flag in flagged["flags"]
            )
            ws.cell(row, 2).comment = Comment(note, "Statement checks")

    # Summary sheet.
    from .engine import summarize

    s = summarize(txns)
    ws2 = wb.create_sheet("Summary")
    rows = [
        ("Metric", "Value"),
        ("Transactions", s.count),
        ("Total debit", _to_float(s.total_debit)),
        ("Total credit", _to_float(s.total_credit)),
        ("Net", _to_float(s.net)),
        ("Currency", s.currency),
        ("Start date", s.start_date or ""),
        ("End date", s.end_date or ""),
    ]
    for r in rows:
        ws2.append(list(r))
    ws2.cell(row=6, column=2).data_type = "s"
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
    for i in (3, 4, 5):
        ws2.cell(row=i, column=2).number_format = money_fmt
    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 18

    _add_checks(wb, txns, header_font, header_fill, money_fmt)

    class MemoryExcelWriter(ExcelWriter):
        """Keep worksheet XML in memory instead of openpyxl's temp files."""

        def write_worksheet(self, worksheet):
            worksheet._drawing = SpreadsheetDrawing()
            worksheet._drawing.charts = worksheet._charts
            worksheet._drawing.images = worksheet._images
            with io.BytesIO() as xml:
                writer = WorksheetWriter(worksheet, out=xml)
                writer.write()
                worksheet._rels = writer._rels
                self._archive.writestr(worksheet.path[1:], writer.read())
                self.manifest.append(worksheet)

    out = io.BytesIO()
    # Workbook.save(BytesIO()) still creates a disk-backed XML file per sheet.
    # This per-conversion writer avoids that without mutating openpyxl globals.
    with ZipFile(out, "w", ZIP_DEFLATED, allowZip64=True) as archive:
        MemoryExcelWriter(wb, archive).save()
    return out.getvalue()
